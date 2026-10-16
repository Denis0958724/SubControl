"""Builds sample bank statements in the formats Ukrainian banks export, for parser tests.
All data is made up. Today is fixed to 2026-10-16 so the tests are stable."""
import csv, datetime as dt, random
import openpyxl, xlwt

OUT = 'src/lib/statement/__tests__/fixtures/'
TODAY = dt.date(2026, 10, 16)
USD = 44.96
random.seed(7)

def monthly(day_offset, n=3):
    d = TODAY + dt.timedelta(days=day_offset)
    out = []
    for _ in range(n):
        out.append(d)
        m = d.month + 1
        y = d.year + (m > 12)
        m = (m - 1) % 12 + 1
        d = d.replace(year=y, month=m)
    return [x for x in out if x <= TODAY]

ops = []  # (date, description, category, amount_uah, op_amount, op_cur)
for d in monthly(-89): ops.append((d, 'Netflix.com', 'Інтернет-сервіси', -round(7.99 * USD, 2), -7.99, 'USD'))
for d in monthly(-86): ops.append((d, 'Spotify', 'Інтернет-сервіси', -149.0, -149.0, 'UAH'))
for d in monthly(-80): ops.append((d, 'APPLE.COM/BILL', 'Інтернет-сервіси', -round(0.99 * USD, 2), -0.99, 'USD'))
for d in monthly(-75): ops.append((d, 'APPLE.COM/BILL', 'Інтернет-сервіси', -round(4.99 * USD, 2), -4.99, 'USD'))
for d in monthly(-72): ops.append((d, 'MEGOGO', 'Інтернет-сервіси', -199.0, -199.0, 'UAH'))
for d in monthly(-70): ops.append((d, 'Спортклуб Sportlife', 'Спорт', -1290.0, -1290.0, 'UAH'))
for d in monthly(-68): ops.append((d, 'Переказ на картку Олена К.', 'Перекази', -2000.0, -2000.0, 'UAH'))
for d in monthly(-60): ops.append((d, 'Поповнення мобільного Київстар', 'Мобільний зв’язок', -150.0, -150.0, 'UAH'))
for i in range(1, 90, 2):
    ops.append((TODAY - dt.timedelta(days=i), f'Сільпо, магазин №{100 + i}', 'Продукти', -float(250 + (i * 37) % 600), None, 'UAH'))
for i in range(2, 90, 5):
    ops.append((TODAY - dt.timedelta(days=i), 'Bolt', 'Таксі', -float(110 + (i * 13) % 120), None, 'UAH'))
for d in monthly(-85): ops.append((d, 'Зарахування зарплати', 'Надходження', 30000.0, 30000.0, 'UAH'))
ops.sort(key=lambda o: o[0], reverse=True)

base_ops = list(ops)

# many distinct strings so the old .xls format has to split its string table across records
for i in range(600):
    ops.append((TODAY - dt.timedelta(days=i % 88 + 1), f'Покупка в інтернет-магазині «Довга назва продавця номер {i}»', 'Покупки', -float(10 + i), None, 'UAH'))
ops.sort(key=lambda o: o[0], reverse=True)

PRIVAT_HEAD = ['Дата', 'Категорія', 'Картка', 'Опис операції', 'Сума в валюті картки', 'Валюта картки',
               'Сума в валюті транзакції', 'Валюта транзакції', 'Залишок на кінець періоду', 'Валюта залишку']

def privat_row(o, as_serial=False):
    d = dt.datetime.combine(o[0], dt.time(14, 30))
    date = d if as_serial else d.strftime('%d.%m.%Y %H:%M:%S')
    op = o[4] if o[4] is not None else o[3]
    return [date, o[2], '5168 **** **** 1234', o[1], o[3], 'UAH', op, o[5], 12345.67, 'UAH']

# Privat24-style .xlsx with a title row, dates as text
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Виписка'
ws.append(['Виписка з Ваших карток за період 18.07.2026 - 16.10.2026'])
ws.append(PRIVAT_HEAD)
for o in ops: ws.append(privat_row(o))
wb.save(OUT + 'privat.xlsx')

# same with real Excel dates
wb = openpyxl.Workbook(); ws = wb.active
ws.append(PRIVAT_HEAD)
for o in ops:
    r = privat_row(o, as_serial=True); ws.append(r)
    ws.cell(row=ws.max_row, column=1).number_format = 'DD.MM.YYYY HH:MM'
wb.save(OUT + 'privat-dates.xlsx')

# old binary .xls
book = xlwt.Workbook(encoding='utf-8'); sh = book.add_sheet('Виписка')
sh.write(0, 0, 'Виписка з Ваших карток за період 18.07.2026 - 16.10.2026')
for c, h in enumerate(PRIVAT_HEAD): sh.write(1, c, h)
for r, o in enumerate(ops, start=2):
    for c, v in enumerate(privat_row(o)): sh.write(r, c, v)
book.save(OUT + 'privat.xls')

# monobank CSV export (UTF-8 with BOM, comma separated)
with open(OUT + 'monobank.csv', 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f)
    w.writerow(['Дата i час операції', 'Деталі операції', 'MCC', 'Сума в валюті картки (UAH)', 'Сума в валюті операції',
                'Валюта', 'Курс', 'Сума комісій (UAH)', 'Сума кешбеку (UAH)', 'Залишок після операції'])
    for o in base_ops:
        mcc = {'Netflix.com': 4899, 'Spotify': 5815, 'APPLE.COM/BILL': 5818, 'MEGOGO': 4899, 'Bolt': 4121}.get(o[1], 5411 if 'Сільпо' in o[1] else 5999)
        op = o[4] if o[4] is not None else o[3]
        w.writerow([dt.datetime.combine(o[0], dt.time(9, 15)).strftime('%d.%m.%Y %H:%M:%S'), o[1], mcc,
                    f'{o[3]:.2f}', f'{op:.2f}', o[5], '—', '0.00', '0.00', '1000.00'])

# a bank that exports Windows-1251 CSV with separate debit / credit columns
with open(OUT + 'other-cp1251.csv', 'w', encoding='cp1251', newline='') as f:
    w = csv.writer(f, delimiter=';')
    w.writerow(['Дата операції', 'Призначення платежу', 'Дебет', 'Кредит', 'Валюта'])
    for o in base_ops:
        amt = o[3]
        w.writerow([o[0].strftime('%d.%m.%Y'), o[1].replace('’', "'"),
                    f'{-amt:.2f}'.replace('.', ',') if amt < 0 else '', f'{amt:.2f}'.replace('.', ',') if amt > 0 else '', 'UAH'])

# HTML table saved as .xls
rows = ''.join(f'<tr><td>{o[0].strftime("%d.%m.%Y")}</td><td>{o[1]}</td><td>{o[3]:.2f}</td><td>UAH</td></tr>' for o in base_ops)
with open(OUT + 'html-table.xls', 'w', encoding='utf-8') as f:
    f.write('<html><head><meta charset="utf-8"></head><body><table>'
            '<tr><th>Дата</th><th>Опис</th><th>Сума</th><th>Валюта</th></tr>' + rows + '</table></body></html>')

print('fixtures:', len(ops), 'operations')
